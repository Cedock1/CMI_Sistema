#!/usr/bin/env python3
"""
Llena la hoja CARGA de `Plantilla_Carga_Inspecciones_GAMLP.xlsx` con los compromisos de un mes,
a partir de las propuestas ya razonadas en `secretos/propuesta_*.json`.

    python3 scripts/llenar_plantilla_gamlp.py --revisar   # muestra las filas, no escribe
    python3 scripts/llenar_plantilla_gamlp.py             # escribe la plantilla (con respaldo)

POR QUÉ EXISTE (César, 17-sep)
  El GAMLP carga las inspecciones en otro sistema, con una plantilla fija. La Base Activa que
  exporta ese sistema llega hasta el 28-jul: agosto no está. Las propuestas del CMI ya tienen
  los compromisos de agosto razonados; este script solo los pasa al formato de la plantilla.
  No llama a ninguna API ni toca la base del CMI.

EL FORMATO, TAL CUAL LO USA LA BASE ACTIVA
  · Una fila por tarea. `Nº INSPECCION` se repite dentro del evento; `Nº TAREA` es correlativo
    global. Los dos siguen la numeración de la Base Activa (agosto arranca en 72 y 296).
  · Texto en MAYÚSCULAS y sin tildes (la Ñ se conserva). INICIO PREVISTO = fecha del evento.
  · Unidad, macrodistrito, estado y dependencia salen EXACTOS de la hoja CATALOGOS.
  · RESPONSABLE: la persona que la Base Activa ya asigna a esa dependencia; si no hay, el
    titular según el reporte de RRHH más reciente. Solo el nombre: nada más de ese archivo.
  · Cada compromiso es una fila, y cada instrucción de la «operativa agrupada» también, porque
    la Base Activa registra instrucciones de ese tamaño («reponer las impresoras…»).
  · Los enriquecimientos NO generan filas: el compromiso ya existe (en la Base Activa o en otra
    fila de este mismo mes).
"""
from __future__ import annotations

import argparse
import collections
import datetime
import json
import re
import shutil
import sys
import unicodedata
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
SECRETOS = RAIZ / "secretos"
# Los archivos del GAMLP salieron de Descargas el 17-sep. `entregas/` NO se versiona: la
# plantilla llena lleva nombres de funcionarios y el .md, las transcripciones íntegras.
ENTREGA = RAIZ / "entregas" / "2026-08 - Inspecciones GAMLP"
PLANTILLA = ENTREGA / "01 - Entregables" / "Plantilla_Carga_Inspecciones_GAMLP.xlsx"
BASE = ENTREGA / "02 - Fuentes" / "GAMLP_Base_Activa_2026-09-17.xlsx"
RESPALDO = ENTREGA / "02 - Fuentes" / "Plantilla_Carga_Inspecciones_GAMLP_vacia.xlsx"
# RRHH trae CI, celular y fecha de nacimiento: vive en `secretos/`, y de acá solo se lee el nombre.
RRHH = SECRETOS / "fuentes" / "Consultor gamlp 07-09-2026.xlsx"
EVENTOS = SECRETOS / "plantilla_agosto_eventos.json"
ESTADO = "NO REPORTADA"   # decisión de César, 17-sep: nadie reportó avance todavía
# Sin plazo en la propuesta: la Base Activa casi nunca deja FIN PREVISTO vacío (5 de 295) y sus
# plazos típicos van de 2 semanas a 3 meses. Operativas → +30 días; compromisos → +90 días.
DIAS_POR_OMISION = {"operativa": 30, "compromiso": 90}

# Responsables escritos como texto libre en las propuestas del 10 al 13-ago → sigla del MOF.
TEXTO_A_SIGLA = [
    ("tutela el laboratorio", "LSM"),
    ("Transparencia / Auditoría", "UTLCC"),
    ("seguridad industrial", "UBSSO"),
    ("Dirección Jurídica", "DAJ"),
    ("Semaforización", "DMUSIT"),
    ("residuos especiales", "DRSECPV"),
    ("Seguridad Ciudadana", "DSCC"),
    ("Recursos Humanos", "DGRH"),
    ("Tributaria", "ATM"),
    ("Alumbrado Público", "USE"),
    ("Dirección de Comunicación", "DC"),
    ("Ciudad Vital", "SMCVI"),
    ("Ciudad Verde", "SMCVE"),
    ("Ciudad Planificada", "SMCPH"),
    ("Despacho", "DAM"),
]

# Entidades y Despacho: el nombre del MOF no se parece al del catálogo.
MANUAL = {
    "DAM": ("SEM GESTION EFICIENTE", "DESPACHO"),
    "SEMGE": ("SEM GESTION EFICIENTE", "SEM GESTION EFICIENTE"),
    # Unidades de control que en el MOF cuelgan de sí mismas: van con el Despacho, como en la Base.
    "UTLCC": ("SEM GESTION EFICIENTE", "U. TRANSPARENCIA Y LUCHA CONTRA LA CORRUPCION"),
    "UAI": ("SEM GESTION EFICIENTE", "U. AUDITORIA INTERNA"),
    # La Base la reparte entre cinco unidades según la empresa; por defecto, el Despacho (MOF).
    "DEESPM": ("SEM GESTION EFICIENTE", "DIR. EMPRESAS, ENTIDADES Y SERVICIOS PUBLICOS MUNICIPALES"),
    "SETRAM": ("SETRAM - SERVICIO DE TRANSPORTE MUNICIPAL",) * 2,
    "UDEPR": ("SM CIUDAD PLANIFICADA Y HABITABLE", "U. DISENOS Y ESTUDIOS DE PREINVERSION"),
    "LSM": ("SM CIUDAD PLANIFICADA Y HABITABLE", "LABORATORIO DE SUELOS Y MATERIALES"),
    "BMVP": ("BMVP - BIOPARQUE MUNICIPAL VESTY PAKOS",) * 2,
    "SIREMU": ("SIREMU - SISTEMA DE REGULACION Y SUPERVISION MUNICIPAL",) * 2,
    "EDMTB": ("EDMTB - ENTIDAD DESCENTRALIZADA MUNICIPAL TERMINAL DE BUSES LA PAZ",) * 2,
    "EMAVERDE": ("EMAVERDE - EMPRESA MUNICIPAL DE AREAS VERDES, PARQUES Y FORESTACION",) * 2,
    "EMAVIAS": ("EMAVIAS - EMPRESA MUNICIPAL DE ASFALTOS Y VIAS",) * 2,
    "EDMC": ("EDMC - ENTIDAD DESCENTRALIZADA MUNICIPAL DE CEMENTERIOS DE LA PAZ",) * 2,
    "EDMME": ("EDMME - ENTIDAD DESCENTRALIZADA MUNICIPAL DE MAQUINARIA Y EQUIPO",) * 2,
    "SAMAPA": ("SERVICIO AUTONOMO MUNICIPAL DE AGUA POTABLE Y ALCANTARILLADO",) * 2,
}
# El puesto de los subalcaldes es «SUB ALCALDE …» y no entra por el patrón de titulares.
SUBALCALDES = {"SAZ": "RUBEN LIMA QUISPE", "SAP": "WALTER CALIXTO ARANDA CHIPANA",
               "SAM": "JANETH FLORES GUZMAN"}


def mayus(texto: str | None) -> str:
    """MAYÚSCULAS sin tildes, conservando la Ñ, como la Base Activa."""
    if not texto:
        return ""
    t = unicodedata.normalize("NFD", str(texto).upper())
    t = "".join(c for c in t if unicodedata.category(c) != "Mn" or c == "̃")
    t = unicodedata.normalize("NFC", t).replace("«", "").replace("»", "").replace("“", "").replace("”", "")
    return re.sub(r"\s+", " ", t).strip()


def clave(t: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[^A-Z0-9Ñ ]", " ", mayus(t))).strip()


def catalogos():
    wb = openpyxl.load_workbook(PLANTILLA)
    filas = list(wb["CATALOGOS"].iter_rows(min_row=2, values_only=True))
    return {i: [f[i] for f in filas if f[i]] for i in range(4)}


def mapa_unidades(cat) -> dict:
    """Sigla del MOF → (unidad, dependencia, responsable) en el vocabulario de la plantilla."""
    sql = (RAIZ / "seed/0002_seed_referencia.sql").read_text(encoding="utf-8")
    mof = {m[1]: dict(nombre=m[2], secretaria=m[3]) for m in re.finditer(
        r"insert into unidad \(id, sigla, nombre, nivel, secretaria[^)]*\) values "
        r"\(\d+, '([^']*)', '([^']*)', '[^']*', '([^']*)'", sql)}
    deps, unidades = cat[3], cat[0]

    def a_catalogo(nombre: str):
        k = clave(nombre)
        k = re.sub(r"^UNIDAD (DE (LA |LOS |LAS )?)?", "U. ", k)
        k = re.sub(r"^DIRECCION (DE (LA |LOS |LAS )?)?", "DIR. ", k)
        k = re.sub(r"^SECRETARIA MUNICIPAL (DE )?", "SM ", k)
        k = re.sub(r"^SUBALCALDIA (DEL MACRODISTRITO )?", "SUB. ", k)
        exacto = [d for d in deps if clave(d) == clave(k)]
        if exacto:
            return exacto[0]
        prefijo = [d for d in deps if clave(k) and clave(d).startswith(clave(k)[:25])]
        return prefijo[0] if prefijo else None

    base = list(openpyxl.load_workbook(BASE)["INSPECCIONES"].iter_rows(min_row=2, values_only=True))
    resp_dep = collections.defaultdict(collections.Counter)
    uni_dep = collections.defaultdict(collections.Counter)
    for f in base:
        if f[8] and f[7] and "\n" not in f[7]:
            resp_dep[f[8]][f[7]] += 1
            uni_dep[f[8]][f[6]] += 1

    titular = {}
    ws = openpyxl.load_workbook(RRHH, read_only=True).worksheets[0]
    it = ws.iter_rows(values_only=True)
    ix = {k: i for i, k in enumerate(next(it))}
    patron = re.compile(r"^(JEFE DE UNIDAD|JEFA DE UNIDAD|JEFE DE LABORATORIO|DIRECTOR|DIRECTORA|SECRETARIO|"
                        r"SECRETARIA MUNICIPAL|SECRETARIA EJECUTIVA|GERENTE|ADMINISTRADOR)")
    for r in it:
        puesto = mayus(r[ix["puesto"]])
        fin = r[ix["fecha_fin"]]
        if (fin and str(fin)[:10] < "2026-09-07") or not patron.match(puesto) or puesto == "SECRETARIA":
            continue
        titular.setdefault(r[ix["sigla"]], mayus(" ".join(
            str(x) for x in (r[ix["nombres"]], r[ix["primer_apellido"]], r[ix["segundo_apellido"]]) if x)))

    out = {}
    for sigla, u in mof.items():
        dep = a_catalogo(u["nombre"])
        sec = mof.get(u["secretaria"], {}).get("nombre")
        uni = a_catalogo(sec) if sec else None
        # Si la Base Activa pone esa dependencia siempre bajo la misma unidad (2+ filas), manda
        # la Base: es donde se va a cargar. Caso real: Educación figura bajo Cuidados y Derechos
        # en la Base y bajo Ciudad Inteligente en el MOF.
        if dep in uni_dep and len(uni_dep[dep]) == 1 and sum(uni_dep[dep].values()) >= 2:
            uni = next(iter(uni_dep[dep]))
        if uni not in unidades:
            uni = uni_dep[dep].most_common(1)[0][0] if dep in uni_dep else None
        if sigla in MANUAL:
            uni, dep = MANUAL[sigla]
        if not uni and u["secretaria"] in ("DAM", "SEMGE"):
            uni = "SEM GESTION EFICIENTE"
        resp = (resp_dep[dep].most_common(1)[0][0] if dep in resp_dep
                else titular.get(sigla) or SUBALCALDES.get(sigla))
        out[sigla] = dict(unidad=uni, dependencia=dep, responsable=resp)
    for sigla, (uni, dep) in MANUAL.items():
        out.setdefault(sigla, dict(unidad=uni, dependencia=dep,
                                   responsable=resp_dep[dep].most_common(1)[0][0] if dep in resp_dep else None))
    out["DAM"]["responsable"] = "CESAR LUIS DOCKWEILER SUAREZ"
    return out


def sigla_de(item: dict, padre: dict | None = None) -> str | None:
    if item.get("responsable_excel"):
        return item["responsable_excel"]
    if item.get("responsable_sigla"):
        return item["responsable_sigla"]
    texto = item.get("responsable_propuesto_texto") or ""
    for patron, sigla in TEXTO_A_SIGLA:
        if patron.lower() in texto.lower():
            return sigla
    return sigla_de(padre) if padre else None


def filas_del_mes(eventos: list, mapa: dict, cat: dict, n_insp: int, n_tarea: int):
    filas, avisos, por_omision = [], [], []
    for ev in eventos:
        d = json.loads((SECRETOS / ev["propuesta"]).read_text(encoding="utf-8"))
        ajustes = ev.get("ajustes", {})
        items = [(it, None) for it in d.get("nuevos", d.get("compromisos", []))]
        op = d.get("operativa_agrupada")
        if op:
            # Sin subtareas, la operativa es en sí misma la instrucción (una visita, un trámite).
            items += [(s, op) for s in op.get("subtareas", [])] or [(op, None)]
        items = [(it, p) for it, p in items if not ajustes.get(it["titulo"], {}).get("omitir")]
        if not items:
            continue
        n_insp += 1
        if ev["macro"] and ev["macro"] not in cat[1]:
            avisos.append(f"macrodistrito fuera de catálogo: {ev['macro']}")
        for it, padre in items:
            it = {**it, **ajustes.get(it["titulo"], {})}
            n_tarea += 1
            sigla = sigla_de(it, padre)
            u = mapa.get(sigla or "", {})
            plazo = it.get("plazo") or (padre or {}).get("plazo")
            if not plazo:
                dias = DIAS_POR_OMISION["operativa" if padre or it.get("eje_sugerido") == "OP" else "compromiso"]
                plazo = (datetime.date.fromisoformat(ev["fecha"]) + datetime.timedelta(days=dias)).isoformat()
                por_omision.append(n_tarea)
            unidad = ev.get("unidad_por_sigla", {}).get(sigla) or it.get("unidad_excel") or u.get("unidad")
            responsable = it.get("responsable_nombre_excel") or u.get("responsable")
            fila = [n_insp, ev["fecha"], mayus(ev["evento"]), n_tarea, mayus(it.get("texto_excel") or it["titulo"]),
                    ev["macro"], unidad, responsable, u.get("dependencia"),
                    ev["fecha"], plazo, None, ESTADO, mayus(it.get("observacion_excel")) or None, None, None]
            if not unidad:
                avisos.append(f"sin unidad ({sigla}): {it['titulo'][:70]}")
            if not u.get("dependencia"):
                avisos.append(f"sin dependencia ({sigla}): {it['titulo'][:70]}")
            elif not responsable:
                avisos.append(f"sin responsable ({sigla} · {u['dependencia']}): {it['titulo'][:60]}")
            if fila[6] and fila[6] not in cat[0]:
                avisos.append(f"unidad fuera de catálogo: {fila[6]}")
            if fila[8] and fila[8] not in cat[3]:
                avisos.append(f"dependencia fuera de catálogo: {fila[8]}")
            filas.append(fila)
    if por_omision:
        avisos.append(f"FIN PREVISTO por omisión en {len(por_omision)} filas (la propuesta no traía plazo)")
    return filas, avisos


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--revisar", action="store_true", help="muestra las filas y no escribe")
    args = ap.parse_args()

    cat = catalogos()
    mapa = mapa_unidades(cat)
    cfg = json.loads(EVENTOS.read_text(encoding="utf-8"))
    base = list(openpyxl.load_workbook(BASE)["INSPECCIONES"].iter_rows(min_row=2, values_only=True))
    ultimo_insp = max(f[0] for f in base if isinstance(f[0], int))
    ultima_tarea = max(f[3] for f in base if isinstance(f[3], int))
    filas, avisos = filas_del_mes(cfg["eventos"], mapa, cat, ultimo_insp, ultima_tarea)

    por_evento = collections.Counter((f[0], f[1], f[2]) for f in filas)
    for (n, fecha, ev), k in por_evento.items():
        print(f"  {n:>3} · {fecha} · {ev[:60]:<60} {k:>3} tareas")
    print(f"\n{len(por_evento)} inspecciones ({ultimo_insp + 1}–{ultimo_insp + len(por_evento)}) · "
          f"{len(filas)} tareas ({ultima_tarea + 1}–{ultima_tarea + len(filas)})")
    for a in sorted(set(avisos)):
        print("  ! " + a)

    if args.revisar:
        for f in filas:
            print(" | ".join("" if v is None else str(v) for v in f[:11]))
        print("\n--revisar: no se escribió nada.")
        return

    RESPALDO.parent.mkdir(parents=True, exist_ok=True)
    if not RESPALDO.exists():
        shutil.copy2(PLANTILLA, RESPALDO)
    wb = openpyxl.load_workbook(RESPALDO)   # siempre desde la plantilla vacía: idempotente
    ws = wb["CARGA"]
    for i, fila in enumerate(filas, start=2):
        for j, v in enumerate(fila, start=1):
            ws.cell(i, j, v)
    wb.save(PLANTILLA)
    print(f"\nEscrita: {PLANTILLA} ({len(filas)} filas). Plantilla vacía respaldada en {RESPALDO}")


if __name__ == "__main__":
    sys.exit(main())
